"""Translator files of ``i18n_export``/``i18n_import``: ``.xlsx`` (openpyxl) or ``.csv`` (UTF-8)."""

from __future__ import annotations

import csv
from pathlib import Path

from django.core.management.base import CommandError

COLUMNS = ["catalog", "context", "source", "source_plural", "translation", "note"]


def write_rows(path: Path, rows: list[dict[str, str]], title: str) -> None:
    if path.suffix.lower() == ".csv":
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=COLUMNS)
            writer.writeheader()
            writer.writerows(rows)
        return
    if path.suffix.lower() != ".xlsx":
        raise CommandError("The file must end in .xlsx or .csv.")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row.get(column, "") for column in COLUMNS])
    for column, width in zip("ABCDEF", (28, 16, 60, 30, 60, 40), strict=True):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    workbook.save(path)


def read_rows(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        raise CommandError(f"[{path}] does not exist.")
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            return [{k: (v or "") for k, v in row.items() if k} for row in csv.DictReader(handle)]
    if path.suffix.lower() != ".xlsx":
        raise CommandError("The file must end in .xlsx or .csv.")
    from openpyxl import load_workbook

    try:
        sheet = load_workbook(path, read_only=True).active
    except Exception as error:  # openpyxl raises several types for broken files
        raise CommandError(f"The file is not a valid spreadsheet: {type(error).__name__}.") from error
    rows = sheet.iter_rows(values_only=True)
    header = [str(value or "") for value in next(rows, [])]
    if "source" not in header or "translation" not in header:
        raise CommandError("The spreadsheet needs 'source' and 'translation' columns.")
    return [
        {name: "" if value is None else str(value) for name, value in zip(header, values, strict=False) if name}
        for values in rows
    ]
