"""Port of tests/Feature/I18n/TranslationCommandsTest.php (gettext edition): i18n_missing,
i18n_export and i18n_import on a throwaway project directory (``--root``)."""

from __future__ import annotations

import csv
import io
from pathlib import Path

import polib
import pytest
from django.core.management import call_command
from openpyxl import load_workbook

pytestmark = pytest.mark.django_db

HEADER = {
    "Content-Type": "text/plain; charset=UTF-8",
    "Language": "ko",
    "Plural-Forms": "nplurals=1; plural=0;",
}


@pytest.fixture
def project(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """An app whose source uses two strings; its Korean catalog has one of them and one the
    source no longer uses (original: lang/en/registration.php vs lang/ko/registration.php)."""
    app = tmp_path / "registration"
    (app / "locale" / "ko" / "LC_MESSAGES").mkdir(parents=True)
    (app / "messages.py").write_text(
        "from django.utils.translation import gettext_lazy as _\n\n"
        'ERRORS = [_("Credit limit exceeded"), _("Closed")]\n',
        encoding="utf-8",
    )
    po = polib.POFile()
    po.metadata = HEADER
    po.append(polib.POEntry(msgid="Closed", msgstr="마감"))
    po.append(polib.POEntry(msgid="Removed", msgstr="삭제됨"))
    po.save(str(app / "locale" / "ko" / "LC_MESSAGES" / "django.po"))
    # i18n_import records its run in the business audit (AuditedCommand).
    monkeypatch.setattr("ninefruits.audit.logger.audit", lambda *args, **kwargs: None)
    return tmp_path


def run(*args: str) -> tuple[int, str]:
    out = io.StringIO()
    code = call_command(*args, stdout=out, stderr=io.StringIO(), no_color=True)
    return int(code or 0), out.getvalue()


def test_reports_missing_and_obsolete_strings(project: Path) -> None:
    code, output = run("i18n_missing", "--strict", "--root", str(project))

    assert code == 2
    assert "1 missing, 1 obsolete" in output
    assert "missing  [registration/locale] Credit limit exceeded" in output
    assert "obsolete [registration/locale] Removed" in output


def test_without_strict_it_only_reports(project: Path) -> None:
    code, _output = run("i18n_missing", "ko", "--root", str(project))
    assert code == 0


def test_fuzzy_empty_and_broken_placeholders_count_as_missing(project: Path) -> None:
    path = project / "registration" / "locale" / "ko" / "LC_MESSAGES" / "django.po"
    po = polib.pofile(str(path))
    po.append(polib.POEntry(msgid="Credit limit exceeded", msgstr="학점 상한 초과", flags=["fuzzy"]))
    po.save()
    (project / "registration" / "more.py").write_text(
        'from django.utils.translation import gettext as _\n\nX = _("%(count)s seats left") % {"count": 1}\n',
        encoding="utf-8",
    )
    po.append(polib.POEntry(msgid="%(count)s seats left", msgstr="%(seats)s석 남음", flags=["python-format"]))
    po.save()

    code, output = run("i18n_missing", "--strict", "--root", str(project))

    assert code == 2
    assert "Credit limit exceeded" in output
    assert "%(count)s seats left (placeholders differ)" in output


def test_strict_passes_when_complete(project: Path) -> None:
    path = project / "registration" / "locale" / "ko" / "LC_MESSAGES" / "django.po"
    po = polib.pofile(str(path))
    po.append(polib.POEntry(msgid="Credit limit exceeded", msgstr="학점 상한 초과"))
    po.save()
    code, output = run("i18n_missing", "--strict", "--root", str(project))
    assert code == 0
    assert "0 missing, 1 obsolete" in output


def test_exports_and_imports_translations_for_translators(project: Path, tmp_path: Path) -> None:
    # The catalog needs the new string first (makemessages does this in the project).
    path = project / "registration" / "locale" / "ko" / "LC_MESSAGES" / "django.po"
    po = polib.pofile(str(path))
    po.append(polib.POEntry(msgid="Credit limit exceeded", msgstr=""))
    po.save()
    workbook = tmp_path / "ko.xlsx"

    code, _output = run("i18n_export", "ko", str(workbook), "--missing", "--root", str(project))

    assert code == 0
    sheet = load_workbook(workbook).active
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    assert rows[0] == ["catalog", "context", "source", "source_plural", "translation", "note"]
    assert [row[:5] for row in rows[1:]] == [["registration/locale", None, "Credit limit exceeded", None, None]]

    sheet.cell(row=2, column=5, value="학점 상한 초과")
    sheet.append(["registration/locale", None, "Unknown string", None, "ignored", None])
    load_target = tmp_path / "ko-translated.xlsx"
    sheet.parent.save(load_target)

    code, output = run("i18n_import", "ko", str(load_target), "--dry-run", "--root", str(project))
    assert code == 0 and "1 translation(s) would be written" in output
    assert polib.pofile(str(path)).find("Credit limit exceeded").msgstr == ""

    code, _output = run("i18n_import", "ko", str(load_target), "--root", str(project))

    assert code == 0
    imported = {e.msgid: e.msgstr for e in polib.pofile(str(path))}
    assert imported == {"Closed": "마감", "Removed": "삭제됨", "Credit limit exceeded": "학점 상한 초과"}
    assert path.with_suffix(".mo").is_file()


def test_csv_round_trip(project: Path, tmp_path: Path) -> None:
    target = tmp_path / "ko.csv"
    run("i18n_export", "ko", str(target), "--root", str(project))
    rows = list(csv.DictReader(target.open(encoding="utf-8-sig")))
    assert {row["source"]: row["translation"] for row in rows} == {"Closed": "마감", "Removed": "삭제됨"}
    rows[0]["translation"] = "접수 마감"
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    code, _output = run("i18n_import", "ko", str(target), "--root", str(project))
    assert code == 0
    path = project / "registration" / "locale" / "ko" / "LC_MESSAGES" / "django.po"
    assert polib.pofile(str(path)).find("Closed").msgstr == "접수 마감"


def test_rejects_invalid_locales(project: Path, tmp_path: Path) -> None:
    from django.core.management.base import CommandError

    with pytest.raises(CommandError, match="Invalid locale"):
        run("i18n_import", "../ko", str(tmp_path / "x.csv"), "--root", str(project))


@pytest.mark.parametrize("app", ["i18n", "settings", "pdf", "institution"])
def test_the_strings_of_the_i18n_work_are_translated(app: str) -> None:
    """Every string of these apps (vendor overrides included) is translated in ko, zh_Hans, lo."""
    from django.conf import settings

    from ninefruits.i18n import catalogs

    extracted = catalogs.extract(Path(settings.BASE_DIR) / "ninefruits" / app)
    problems = []
    for locale_dir, keys in extracted.messages.items():
        for locale in catalogs.TRANSLATED_LOCALES:
            entries = {catalogs.entry_key(e): e for e in catalogs.open_po(catalogs.po_path(locale_dir, locale))}
            for key in keys:
                entry = entries.get(key)
                if (
                    entry is None
                    or entry.obsolete
                    or not catalogs.is_translated(entry)
                    or catalogs.format_mismatch(entry)
                ):
                    problems.append((locale, key))
    assert problems == []
